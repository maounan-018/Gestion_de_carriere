#!/usr/bin/env python3
"""
services/profil_import_service.py — Import de profils depuis un classeur Excel.

RÔLE
    Importe des profils (collaborateurs) depuis un .xlsx traité 100 % LOCALEMENT
    (openpyxl, jamais pandas — cohérent avec le socle). 1re ligne = en-têtes ;
    colonnes reconnues (accents/casse tolérés) :
      nom (obligatoire) · prenom · matricule · poste_actuel · departement_key ·
      competences (séparées par des virgules → ProfilCompetence, source=declaratif).
    Import robuste : une ligne fautive n'interrompt pas les autres (collectée dans
    `erreurs`). Données importées = RÉELLES (synthetique=False), pas des personas.

EMPLACEMENT MVC
    Couche "services" : parsing + persistance. Reçoit `db` en paramètre.
"""

from __future__ import annotations

import io
import logging

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.profil import Profil, ProfilCompetence, CompetenceSource

logger = logging.getLogger(__name__)

_MAX_ROWS = 5000  # garde-fou (au-delà = découper le fichier)

# En-tête (normalisé) -> champ du modèle.
_ALIASES = {
    "nom": "nom",
    "prenom": "prenom", "prénom": "prenom",
    "matricule": "matricule",
    "poste": "poste_actuel", "poste_actuel": "poste_actuel", "poste actuel": "poste_actuel",
    "departement": "departement_key", "departement_key": "departement_key",
    "département": "departement_key", "pole": "departement_key", "pôle": "departement_key",
    "competences": "competences", "compétences": "competences",
}


def _norm(h) -> str:
    return str(h or "").strip().lower()


def importer_profils(db: Session, data: bytes) -> dict:
    """Importe des profils depuis les octets d'un .xlsx. Renvoie
    {importes, ignorees, erreurs[]}. Lève ValueError si le fichier est illisible
    ou si la colonne 'nom' manque."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — fichier non-xlsx / corrompu
        raise ValueError(f"Fichier Excel illisible : {e}")
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return {"importes": 0, "ignorees": 0, "erreurs": ["Fichier vide."]}

    cols: dict[int, str] = {}
    for i, h in enumerate(header):
        champ = _ALIASES.get(_norm(h))
        if champ:
            cols[i] = champ
    if "nom" not in cols.values():
        raise ValueError("Colonne « nom » manquante dans l'en-tête.")

    importes = 0
    ignorees = 0
    erreurs: list[str] = []
    for n, row in enumerate(rows, start=2):
        if n - 1 > _MAX_ROWS:
            erreurs.append(f"Import limité à {_MAX_ROWS} lignes ; le reste a été ignoré.")
            break
        vals = {champ: (str(row[i]).strip() if i < len(row) and row[i] is not None else "")
                for i, champ in cols.items()}
        nom = vals.get("nom", "").strip()
        if not nom:
            ignorees += 1
            continue
        try:
            # Savepoint par ligne : une ligne fautive ne défait QUE sa propre insertion.
            with db.begin_nested():
                p = Profil(
                    nom=nom, prenom=vals.get("prenom") or None, matricule=vals.get("matricule") or None,
                    poste_actuel=vals.get("poste_actuel") or None,
                    departement_key=vals.get("departement_key") or None,
                    synthetique=False,  # import de données RÉELLES traitées localement
                )
                db.add(p)
                db.flush()
                for lib in (vals.get("competences") or "").split(","):
                    lib = lib.strip()
                    if lib:
                        db.add(ProfilCompetence(profil_id=p.id, libelle=lib,
                                                source=CompetenceSource.DECLARATIF))
            importes += 1
        except Exception as e:  # noqa: BLE001 — une ligne fautive n'arrête pas l'import
            erreurs.append(f"Ligne {n} ignorée : {e}")

    db.commit()
    logger.info("Import Excel : %d profils importés, %d ignorés, %d erreur(s).",
                importes, ignorees, len(erreurs))
    return {"importes": importes, "ignorees": ignorees, "erreurs": erreurs}
