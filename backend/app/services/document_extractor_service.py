#!/usr/bin/env python3
"""
services/document_extractor_service.py — Extraction de TEXTE depuis un fichier.

RÔLE
    Convertit le contenu binaire d'un fichier téléversé (PDF, Word, Excel, texte)
    en texte brut prêt à être découpé puis indexé (RAG du Coach). Le type est
    déduit de l'extension du nom de fichier. Les libs lourdes sont importées à la
    demande pour ne pas pénaliser le boot.

PIPELINE PDF
    1) Extraction texte native via PyMuPDF (bien plus robuste que pypdf sur les
       mises en page complexes). 2) Si le résultat est « pauvre » (quasi vide,
       squelette de puces…), repli automatique sur l'OCR local rapidocr-onnxruntime
       (CNDP : aucune donnée ne sort, modèles ONNX locaux). 3) Garde-fou final :
       un contenu toujours indigent (PDF image sans texte lisible même par l'OCR)
       lève ValueError → le controller remonte un 400 clair plutôt que d'indexer
       du vide (le Coach répondrait alors du générique sans savoir pourquoi).

SÉCURITÉ
    100 % local. Le texte extrait est traité comme DONNÉE non fiable par
    l'appelant (fencé dans les prompts coach via ai_router.fence).

EMPLACEMENT MVC
    Couche "services" : utilitaire métier, sans route HTTP.
"""

from __future__ import annotations

import io
import logging
import re

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".xlsx", ".txt", ".md")

# Détection d'extraction appauvrie : un run de 3+ lettres = un mot ; un run de
# 2+ lettres sur une ligne = la preuve que cette ligne porte du texte.
_WORD_RE = re.compile(r"[^\W\d_]{3,}", re.UNICODE)
_HAS_WORD_RE = re.compile(r"[^\W\d_]{2,}", re.UNICODE)


def file_extension(filename: str) -> str:
    """Extension en minuscules (point inclus), '' si absente."""
    name = (filename or "").lower().strip()
    dot = name.rfind(".")
    return name[dot:] if dot != -1 else ""


def extract_text(filename: str, content: bytes) -> str:
    """Extrait le texte d'un fichier selon son extension.

    Lève ValueError si le format n'est pas supporté, si rien n'est extractible
    ou si l'extraction est si pauvre qu'elle ne nourrirait pas le RAG utilement.
    """
    ext = file_extension(filename)
    if ext == ".pdf":
        text = _extract_pdf(content)
    elif ext == ".docx":
        text = _extract_docx(content)
    elif ext == ".xlsx":
        text = _extract_xlsx(content)
    elif ext in (".txt", ".md"):
        text = _extract_plain(content)
    else:
        raise ValueError(
            f"Format non supporté : '{ext or '?'}'. "
            f"Formats acceptés : {', '.join(SUPPORTED_EXTENSIONS)}."
        )

    text = (text or "").strip()
    if not text:
        raise ValueError("Aucun texte extractible de ce fichier.")
    if _looks_poorly_extracted(text):
        raise ValueError(
            "Le document a été lu mais son contenu semble illisible ou mal "
            "structuré (trop peu de texte exploitable — PDF scanné/en image ou "
            "mise en page complexe). Fournissez un PDF contenant du vrai texte, "
            "ou une version .docx / .txt du document."
        )
    return text


# --------------------------------------------------------------- Garde-fou qualité

def _looks_poorly_extracted(text: str) -> bool:
    """Détecte une extraction appauvrie : document lu mais vidé de son contenu.

    Cas typiques : PDF scanné où l'OCR a échoué, PDF infographique dont seul le
    squelette (titres + puces) a été capté. Indexer un tel contenu donne au RAG
    un contexte inutile (le Coach répondrait alors de façon générique sans
    savoir pourquoi). On préfère refuser l'import avec un message clair.
    """
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    if not lines:
        return True
    # Structure effondrée : la majorité des lignes ne portent aucun mot
    # (puces « • », coches, séparateurs, caractères isolés).
    symbol_lines = sum(1 for ln in lines if not _HAS_WORD_RE.search(ln))
    if len(lines) >= 8 and symbol_lines / len(lines) >= 0.5:
        return True
    # Contenu non négligeable mais quasi aucun mot réel (ligatures/garbage OCR).
    if len("".join(text.split())) > 40 and len(_WORD_RE.findall(text)) < 5:
        return True
    return False


# --------------------------------------------------------------- Extracteurs

def _extract_plain(content: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def _extract_docx(content: bytes) -> str:
    import docx  # python-docx
    document = docx.Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _extract_xlsx(content: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"# Feuille : {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def _extract_pdf(content: bytes) -> str:
    # PyMuPDF (fitz) : extraction native robuste sur les mises en page complexes.
    try:
        import pymupdf  # PyMuPDF >= 1.24
    except ImportError:  # pragma: no cover
        import fitz as pymupdf

    with pymupdf.open(stream=content, filetype="pdf") as doc:
        text = "\n\n".join((page.get_text("text") or "") for page in doc)
        # PDF en images (scan, guide infographique) : couche texte vide → OCR local.
        if _looks_poorly_extracted(text):
            ocr_text = _ocr_pdf_pages(doc)
            if len(ocr_text.strip()) > len(text.strip()):
                return ocr_text
    return text


# Moteur OCR chargé paresseusement puis réutilisé (chargement ONNX ~1-2 s).
_ocr_engine = None
_OCR_DPI = 200          # compromis qualité/vitesse
_OCR_MAX_PAGES = 50     # borne anti-emballement sur un très gros PDF


def _get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR  # import tardif : lourd (onnx)
        _ocr_engine = RapidOCR()
    return _ocr_engine


def _ocr_pdf_pages(doc) -> str:
    """OCR local des pages (texte dans des images). CNDP : 100 % local/hors-ligne.
    Si rapidocr est absent, renvoie "" → l'appelant garde le texte natif (le
    garde-fou d'extract_text rejettera alors proprement l'import)."""
    try:
        engine = _get_ocr_engine()
    except ImportError:
        logger.warning("OCR indisponible : rapidocr-onnxruntime n'est pas installé.")
        return ""

    n_pages = min(doc.page_count, _OCR_MAX_PAGES)
    if doc.page_count > _OCR_MAX_PAGES:
        logger.warning("OCR limité aux %d premières pages (document de %d).",
                       _OCR_MAX_PAGES, doc.page_count)
    parts: list[str] = []
    for i in range(n_pages):
        try:
            pixmap = doc[i].get_pixmap(dpi=_OCR_DPI)
            result, _ = engine(pixmap.tobytes("png"))
        except Exception as exc:  # noqa: BLE001 — une page illisible n'annule pas tout
            logger.warning("OCR page %d échouée : %s", i + 1, exc)
            continue
        if result:
            parts.append("\n".join(line[1] for line in result))
    text = "\n\n".join(parts)
    logger.info("OCR terminé : %d page(s), %d caractères.", n_pages, len(text))
    return text
