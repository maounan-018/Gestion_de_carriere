#!/usr/bin/env python3
"""
services/referentiels/sfia_loader.py — Ingestion du référentiel SFIA 9 (IT/digital, FR+EN).

RÔLE
    Charge SFIA 9 (tableur Excel officiel, feuille « Skills »/« Compétences ») dans
    le graphe : ~147 compétences numériques bilingues, jointure FR↔EN par le CODE
    SFIA (langue-indépendant : PROG, ARCH, ITSP…). Chaque compétence est tracée
    dans le crosswalk (SkillMapping, source=SFIA, external_id=code) et typée DIGITAL,
    avec son niveau d'entrée SFIA (1..7). Pas de métiers (SFIA = compétences seules).

    Idempotent : ne recharge pas si SFIA déjà présent. Vectorisation faite à part
    (graph_builder), comme pour ESCO/O*NET.

SÉCURITÉ / LICENCE
    Données SFIA sous licence SFIA Foundation (gratuite pour l'usage interne de
    gestion des personnes). Mention d'attribution obligatoire posée sur la source
    (voir core/attributions). 100 % local.

EMPLACEMENT MVC
    Couche "services" : logique d'ingestion, sans route HTTP.
"""

import os
import glob
import logging
from datetime import datetime, timezone

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core import attributions
from app.models.referential import ReferentialSource, SourceCode
from app.models.skill import Skill, SkillType
from app.models.skill_mapping import SkillMapping, MappingMethod

logger = logging.getLogger(__name__)

SFIA_VERSION = "9"

# Index de colonnes (0-based) de la feuille Skills — cf. structure du tableur SFIA 9.
_C_LEVELS = slice(1, 8)   # colonnes 2..8 : niveaux d'application (valeurs 1..7, éparses)
_C_CODE = 8               # colonne 9  : Code SFIA (clé de jointure FR/EN)
_C_SKILL = 10             # colonne 11 : Skill (libellé)
_C_DESC = 13              # colonne 14 : Overall description


def _find_xlsx(lang_token: str, exclude: str | None = None) -> str:
    """Localise le .xlsx SFIA d'une langue (par un jeton du nom de fichier),
    sans dépendre de l'encodage du nom de dossier (« Français »)."""
    base = settings.REFERENTIELS_DIR
    for f in glob.glob(os.path.join(base, "SFIA 9 Excel*", "*.xlsx")):
        name = os.path.basename(f)
        if lang_token in name and (exclude is None or exclude not in name):
            return f
    raise FileNotFoundError(f"Fichier SFIA ({lang_token}) introuvable sous {base}.")


def _skill_rows(path: str) -> list[tuple]:
    """Renvoie les lignes de données de la feuille des compétences (hors en-tête)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next((s for s in ("Skills", "Compétences") if s in wb.sheetnames), wb.sheetnames[0])
    rows = list(wb[sheet].iter_rows(min_row=2, values_only=True))
    wb.close()
    return rows


def _min_level(row: tuple) -> int | None:
    """Niveau SFIA le plus bas auquel la compétence est définie (colonnes 2..8)."""
    levels = [int(c) for c in row[_C_LEVELS] if isinstance(c, (int, float))]
    return min(levels) if levels else None


def _clean(value) -> str:
    return (str(value).replace("_x000D_", " ").strip() if value is not None else "")


def load_sfia(db: Session) -> dict:
    """Ingestion SFIA 9 → graphe. Renvoie un dict de comptes."""
    if db.scalar(select(ReferentialSource).where(ReferentialSource.code == SourceCode.SFIA)):
        logger.info("SFIA déjà ingéré — saut.")
        return {"skipped": True}

    en_path = _find_xlsx("_en_", exclude="_en_US")   # variante britannique, pas l'US
    fr_path = _find_xlsx("_fr_")
    logger.info("Ingestion SFIA — EN=%s | FR=%s", os.path.basename(en_path), os.path.basename(fr_path))

    fr_by_code = {}
    for r in _skill_rows(fr_path):
        code = _clean(r[_C_CODE])
        if code:
            fr_by_code[code] = r

    made: list[tuple[Skill, str, str]] = []   # (skill, code, label_en)
    for r in _skill_rows(en_path):
        code = _clean(r[_C_CODE])
        label_en = _clean(r[_C_SKILL])
        if not code or not label_en:
            continue
        rf = fr_by_code.get(code)
        label_fr = (_clean(rf[_C_SKILL]) if rf else "") or label_en
        desc_en = _clean(r[_C_DESC]) or None
        desc_fr = (_clean(rf[_C_DESC]) if rf else "") or None
        skill = Skill(
            label_fr=label_fr[:300], label_en=label_en[:300],
            description_fr=desc_fr, description_en=desc_en,
            skill_type=SkillType.DIGITAL, sfia_level=_min_level(r),
        )
        db.add(skill)
        made.append((skill, code, label_en))

    db.flush()  # attribue les ids
    db.bulk_insert_mappings(SkillMapping, [
        {"skill_id": s.id, "source": SourceCode.SFIA, "external_id": code,
         "external_label": label_en, "confidence": 1.0,
         "method": MappingMethod.EXACT, "validated": True}
        for (s, code, label_en) in made
    ])

    db.add(ReferentialSource(
        code=SourceCode.SFIA, name="SFIA 9 (SFIA Foundation)", version=SFIA_VERSION,
        languages="fr,en", license="SFIA Foundation (gratuite — usage interne)",
        attribution=attributions.SFIA, imported_at=datetime.now(timezone.utc),
        n_skills=len(made), n_occupations=0,
    ))
    db.commit()

    logger.info("SFIA ingéré : %d compétences numériques (FR/EN).", len(made))
    return {"skills": len(made)}
